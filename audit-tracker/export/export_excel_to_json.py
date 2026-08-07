import json
import sys
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
SRC = sys.argv[1] if len(sys.argv) > 1 else str(HERE.parent / "稽核建議辦理情形追蹤表_範本.xlsx")
OUT = sys.argv[2] if len(sys.argv) > 2 else str(HERE / "rows.json")

wb = openpyxl.load_workbook(SRC)
ws = wb["稽核追蹤表"]

HEADER_ROW = 3
rows = []
for r in range(HEADER_ROW + 1, ws.max_row + 1):
    no = ws.cell(row=r, column=1).value
    unit = ws.cell(row=r, column=2).value
    if not no and not unit:
        continue
    if not no:
        # reserved blank row for future use - skip in the filing export
        continue
    rows.append({
        "no": no,
        "unit": unit or "",
        "item": ws.cell(row=r, column=3).value or "",
        "period": ws.cell(row=r, column=4).value or "",
        "deficiency": ws.cell(row=r, column=5).value or "",
        "action": ws.cell(row=r, column=6).value or "",
        "result": ws.cell(row=r, column=7).value or "",
        "auditor": ws.cell(row=r, column=8).value or "",
        "doc_no": ws.cell(row=r, column=9).value or "",
        "report_date": ws.cell(row=r, column=10).value or "",
    })

title_cell = ws["A1"].value or ""

with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"title": title_cell, "rows": rows}, f, ensure_ascii=False, indent=2)

print(f"exported {len(rows)} rows to {OUT}")
